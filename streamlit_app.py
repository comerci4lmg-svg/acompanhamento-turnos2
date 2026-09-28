from __future__ import annotations

import calendar
from datetime import date, datetime, timedelta
from io import BytesIO
import re
import unicodedata
from zoneinfo import ZoneInfo

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


st.set_page_config(page_title="Acompanhamento de Turnos GO", page_icon="📊", layout="wide")
GRUPOS = ["GOOL", "GOOC", "GOOK", "GOOH"]
EQUIPES_DESMOBILIZADAS = {
    "GOOH013M",
    "GOOL007M",
    "GOOL021M",
    "GOOL024M",
    "GOOL025M",
    "GOOK013M",
    "GOOK012M",
    "GOOK010M",
}
MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]
FUSO_GOIAS = ZoneInfo("America/Sao_Paulo")
META_HORAS_PADRAO = 8.0
META_HORAS_GOOH = 9.0
META_HORAS_GOOH_DS = 8.0


def normalizar_nome(valor):
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    return re.sub(r"[^A-Z0-9]+", "_", texto.upper()).strip("_")


def converter_data_hora(valores):
    """Aceita o ISO enviado pelo bot e datas brasileiras digitadas na planilha."""
    texto = valores.astype("string").str.strip()
    resultado = pd.Series(pd.NaT, index=valores.index, dtype="datetime64[ns]")
    formato_iso = texto.str.match(r"^\d{4}-\d{2}-\d{2}", na=False)
    resultado.loc[formato_iso] = pd.to_datetime(
        texto.loc[formato_iso], errors="coerce", format="ISO8601"
    )
    resultado.loc[~formato_iso] = pd.to_datetime(
        texto.loc[~formato_iso], errors="coerce", format="mixed", dayfirst=True
    )
    return resultado


def formatar_duracao(valor):
    if pd.isna(valor):
        return ""
    if not isinstance(valor, (int, float)):
        return str(valor)
    minutos_totais = round(abs(float(valor)) * 60)
    horas, minutos = divmod(minutos_totais, 60)
    sinal = "-" if float(valor) < 0 else ""
    return f"{sinal}{horas:02d}:{minutos:02d}"


def primeiro_valor_preenchido(valores, padrao=""):
    for valor in valores:
        if not pd.isna(valor) and str(valor).strip():
            return valor
    return padrao


def juntar_valores_distintos(valores):
    vistos = []
    for valor in valores:
        if pd.isna(valor):
            continue
        texto = str(valor).strip()
        if texto and texto not in vistos:
            vistos.append(texto)
    return " | ".join(vistos)


def consolidar_turnos_por_equipe_dia(dados):
    """Une reaberturas da mesma equipe no dia e soma os intervalos oficiais."""
    dados_validos = dados[dados["INICIO_TURNO"].notna()].copy()
    linhas = []
    for (data_turno, grupo, prefixo), registros in dados_validos.groupby(
        ["DATA", "GRUPO", "PREFIXO"], sort=True
    ):
        registros = registros.sort_values("INICIO_TURNO").reset_index(drop=True)
        inicio = registros.iloc[0]["INICIO_TURNO"]
        ultimo = registros.iloc[-1]
        fim = ultimo["FIM_TURNO"]
        em_andamento = pd.isna(fim)

        if "INTERVALO_OFICIAL_SEGUNDOS" in registros.columns:
            intervalo_segundos = float(
                pd.to_numeric(registros["INTERVALO_OFICIAL_SEGUNDOS"], errors="coerce")
                .fillna(0)
                .sum()
            )
        else:
            intervalo_segundos = 0.0
            fim_cobertura = None
            for _, registro in registros.iterrows():
                inicio_registro = registro["INICIO_TURNO"]
                fim_registro = registro["FIM_TURNO"]
                if fim_cobertura is not None and inicio_registro > fim_cobertura:
                    intervalo_segundos += (inicio_registro - fim_cobertura).total_seconds()
                if not pd.isna(fim_registro):
                    if fim_cobertura is None or fim_registro > fim_cobertura:
                        fim_cobertura = fim_registro

        intervalo_horas = intervalo_segundos / 3600
        permanencia_horas = float("nan")
        trabalho_horas = float("nan")
        if not em_andamento:
            permanencia_horas = max(0.0, (fim - inicio).total_seconds() / 3600)
            # Para o acompanhamento, a jornada vai da primeira abertura ao último
            # fechamento. Os intervalos são exibidos separadamente, sem desconto.
            trabalho_horas = permanencia_horas

        saida_prevista = registros.iloc[0]["SAIDA_PREVISTA"]
        diferenca_fechamento = pd.NA
        if not em_andamento and not pd.isna(saida_prevista):
            diferenca_fechamento = round((fim - saida_prevista).total_seconds() / 60)

        linhas.append({
            "DATA": data_turno,
            "GRUPO": grupo,
            "PREFIXO": prefixo,
            "INICIO_TURNO": inicio,
            "SAIDA_PREVISTA": saida_prevista,
            "FIM_TURNO": pd.NaT if em_andamento else fim,
            "SITUACAO": "EM ANDAMENTO" if em_andamento else "ENCERRADO",
            "ABERTURAS_NO_DIA": registros["HIST_TURMA_PLANTAO_ID"].nunique(),
            "PERMANENCIA_HORAS": round(permanencia_horas, 4),
            "INTERVALO_HORAS": round(intervalo_horas, 4),
            "DURACAO_HORAS": round(trabalho_horas, 4),
            "PERMANENCIA": "EM CURSO" if em_andamento else formatar_duracao(permanencia_horas),
            "INTERVALO": formatar_duracao(intervalo_horas),
            "DURACAO": "EM CURSO" if em_andamento else formatar_duracao(trabalho_horas),
            "DIFERENCA_FECHAMENTO_MIN": diferenca_fechamento,
            "PARTICIPA_ESCALA": primeiro_valor_preenchido(registros["PARTICIPA_ESCALA"]),
            "OBSERVACAO": juntar_valores_distintos(registros["OBSERVACAO"]),
            "MOTIVOS_INTERVALO": juntar_valores_distintos(registros["MOTIVOS_INTERVALO"])
                if "MOTIVOS_INTERVALO" in registros.columns else "",
            "HIST_TURMA_PLANTAO_ID": f"{prefixo}-{data_turno}",
        })
    return pd.DataFrame(linhas)


@st.cache_data(ttl=60, show_spinner=False)
def carregar_turnos():
    configuracao = st.secrets["apps_script"]
    resposta = requests.post(
        configuracao["url"],
        json={"action": "read_turnos", "token": configuracao["read_token"]},
        timeout=60,
    )
    resposta.raise_for_status()
    payload = resposta.json()
    if not payload.get("ok"):
        raise RuntimeError(payload.get("error", "Apps Script recusou a leitura."))
    dados = pd.DataFrame(payload.get("registros", []), columns=payload.get("colunas", []))
    return dados, payload.get("atualizadoEm", "")


def preparar_dados(brutos):
    dados = brutos.copy()
    dados.columns = [normalizar_nome(coluna) for coluna in dados.columns]
    obrigatorias = {"PREFIXO", "INICIO_TURNO"}
    ausentes = sorted(obrigatorias - set(dados.columns))
    if ausentes:
        raise ValueError("Colunas ausentes na planilha: " + ", ".join(ausentes))

    for coluna in [
        "INICIO_TURNO", "SAIDA_PREVISTA", "FIM_TURNO",
        "INICIO_INTERVALO", "FIM_INTERVALO",
    ]:
        if coluna not in dados.columns:
            dados[coluna] = pd.NaT
        dados[coluna] = converter_data_hora(dados[coluna])

    dados["PREFIXO"] = dados["PREFIXO"].astype(str).str.strip().str.upper()
    dados["GRUPO"] = dados["PREFIXO"].str[:4]
    dados["DATA"] = dados["INICIO_TURNO"].dt.date
    dados = dados[dados["GRUPO"].isin(GRUPOS)].copy()
    for coluna in ["PARTICIPA_ESCALA", "OBSERVACAO"]:
        if coluna not in dados.columns:
            dados[coluna] = ""
    if "INTERVALO_OFICIAL_SEGUNDOS" in dados.columns:
        dados["INTERVALO_OFICIAL_SEGUNDOS"] = pd.to_numeric(
            dados["INTERVALO_OFICIAL_SEGUNDOS"], errors="coerce"
        ).fillna(0)
    if "MOTIVOS_INTERVALO" not in dados.columns:
        if "MOTIVO_INTERVALO" in dados.columns:
            dados["MOTIVOS_INTERVALO"] = dados["MOTIVO_INTERVALO"]
        else:
            dados["MOTIVOS_INTERVALO"] = ""
    if "HIST_TURMA_PLANTAO_ID" not in dados.columns:
        dados["HIST_TURMA_PLANTAO_ID"] = range(1, len(dados) + 1)
    consolidados = consolidar_turnos_por_equipe_dia(dados)
    if consolidados.empty:
        return consolidados
    consolidados["DIFERENCA_FECHAMENTO_MIN"] = pd.array(
        consolidados["DIFERENCA_FECHAMENTO_MIN"], dtype="Int64"
    )
    return consolidados.sort_values(["INICIO_TURNO", "PREFIXO"], na_position="last")


def preparar_intervalos_individuais(brutos):
    """Prepara um registro por intervalo oficial para o ranking."""
    dados = brutos.copy()
    dados.columns = [normalizar_nome(coluna) for coluna in dados.columns]
    necessarias = {
        "INTERVALO_ID", "PREFIXO", "INICIO_INTERVALO", "FIM_INTERVALO",
        "INTERVALO_OFICIAL_SEGUNDOS",
    }
    if not necessarias.issubset(dados.columns):
        return pd.DataFrame()

    dados["INICIO_INTERVALO"] = converter_data_hora(dados["INICIO_INTERVALO"])
    dados["FIM_INTERVALO"] = converter_data_hora(dados["FIM_INTERVALO"])
    dados["INTERVALO_OFICIAL_SEGUNDOS"] = pd.to_numeric(
        dados["INTERVALO_OFICIAL_SEGUNDOS"], errors="coerce"
    )
    dados["PREFIXO"] = dados["PREFIXO"].astype(str).str.strip().str.upper()
    dados["GRUPO"] = dados["PREFIXO"].str[:4]
    if "MOTIVO_INTERVALO" not in dados.columns:
        dados["MOTIVO_INTERVALO"] = ""

    validos = (
        dados["INTERVALO_ID"].notna()
        & dados["INTERVALO_ID"].astype(str).str.strip().ne("")
        & dados["INICIO_INTERVALO"].notna()
        & dados["FIM_INTERVALO"].notna()
        & dados["INTERVALO_OFICIAL_SEGUNDOS"].ge(0)
        & dados["GRUPO"].isin(GRUPOS)
    )
    intervalos = dados.loc[validos].drop_duplicates("INTERVALO_ID").copy()
    intervalos["INTERVALO_HORAS"] = intervalos["INTERVALO_OFICIAL_SEGUNDOS"] / 3600
    intervalos["DURACAO_INTERVALO"] = intervalos["INTERVALO_HORAS"].map(formatar_duracao)
    return intervalos


def domingo_de_pascoa(ano):
    """Calcula a Páscoa pelo algoritmo gregoriano de Meeus/Jones/Butcher."""
    a = ano % 19
    b = ano // 100
    c = ano % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return date(ano, mes, dia)


def feriados_brasil(ano):
    pascoa = domingo_de_pascoa(ano)
    return {
        date(ano, 1, 1),   # Confraternização Universal
        pascoa - timedelta(days=2),  # Sexta-feira Santa
        date(ano, 4, 21),  # Tiradentes
        date(ano, 5, 1),   # Dia do Trabalho
        pascoa + timedelta(days=60),  # Corpus Christi
        date(ano, 9, 7),   # Independência
        date(ano, 10, 12), # Nossa Senhora Aparecida
        date(ano, 11, 2),  # Finados
        date(ano, 11, 15), # Proclamação da República
        date(ano, 11, 20), # Consciência Negra
        date(ano, 12, 25), # Natal
    }


def ocultar_desmobilizadas_sem_movimento(dados, ano, mes):
    datas = pd.to_datetime(dados["DATA"], errors="coerce")
    equipes_com_movimento = set(
        dados.loc[datas.dt.year.eq(ano) & datas.dt.month.eq(mes), "PREFIXO"]
        .dropna()
        .astype(str)
    )
    manter = (
        ~dados["PREFIXO"].isin(EQUIPES_DESMOBILIZADAS)
        | dados["PREFIXO"].isin(equipes_com_movimento)
    )
    return dados[manter].copy()


def montar_mapa_mensal(dados, ano, mes, grupos, equipes_selecionadas):
    quantidade_dias = calendar.monthrange(ano, mes)[1]
    dias = list(range(1, quantidade_dias + 1))
    hoje = datetime.now(FUSO_GOIAS).date()
    feriados = feriados_brasil(ano)

    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes_selecionadas:
        universo = universo[universo["PREFIXO"].isin(equipes_selecionadas)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    equipes_mapa = sorted(universo["PREFIXO"].dropna().unique())

    datas = pd.to_datetime(universo["DATA"], errors="coerce")
    aberturas_mes = universo[datas.dt.year.eq(ano) & datas.dt.month.eq(mes)].copy()
    aberturas_mes["DIA"] = pd.to_datetime(aberturas_mes["DATA"]).dt.day
    presencas = set(zip(aberturas_mes["PREFIXO"], aberturas_mes["DIA"]))

    linhas = []
    for equipe in equipes_mapa:
        linha = {"EQUIPE": equipe}
        total = 0
        for dia in dias:
            data_celula = date(ano, mes, dia)
            if (equipe, dia) in presencas:
                valor = "1"
                total += 1
            elif data_celula > hoje:
                valor = ""
            elif data_celula in feriados:
                valor = "F"
            elif data_celula.weekday() == 5:
                valor = "S"
            elif data_celula.weekday() == 6:
                valor = "D"
            else:
                valor = "0"
            linha[dia] = valor
        linha["TOTAL"] = total
        linhas.append(linha)
    return pd.DataFrame(linhas).set_index("EQUIPE") if linhas else pd.DataFrame()


def estilo_mapa(valor):
    estilos = {
        "1": "background-color: #16a34a; color: white; font-weight: 700;",
        "0": "background-color: #dc2626; color: white; font-weight: 700;",
        "S": "background-color: #dbeafe; color: #1e3a8a; font-weight: 700;",
        "D": "background-color: #e2e8f0; color: #334155; font-weight: 700;",
        "F": "background-color: #fef3c7; color: #92400e; font-weight: 700;",
        "": "background-color: #f8fafc; color: #94a3b8;",
    }
    return estilos.get(str(valor), "font-weight: 700;")


def montar_tabela_horas(dados, ano, mes, grupos, equipes_selecionadas):
    quantidade_dias = calendar.monthrange(ano, mes)[1]
    dias = list(range(1, quantidade_dias + 1))
    hoje = datetime.now(FUSO_GOIAS).date()
    feriados = feriados_brasil(ano)

    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes_selecionadas:
        universo = universo[universo["PREFIXO"].isin(equipes_selecionadas)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    equipes_tabela = sorted(universo["PREFIXO"].dropna().unique())

    datas = pd.to_datetime(universo["DATA"], errors="coerce")
    turnos_mes = universo[datas.dt.year.eq(ano) & datas.dt.month.eq(mes)].copy()
    turnos_mes["DIA"] = pd.to_datetime(turnos_mes["DATA"]).dt.day

    linhas = []
    for equipe in equipes_tabela:
        linha = {"EQUIPE": equipe}
        total_horas = 0.0
        turnos_equipe = turnos_mes[turnos_mes["PREFIXO"].eq(equipe)]
        for dia in dias:
            data_celula = date(ano, mes, dia)
            turnos_dia = turnos_equipe[turnos_equipe["DIA"].eq(dia)]
            if not turnos_dia.empty:
                if turnos_dia["FIM_TURNO"].isna().any():
                    valor = "EM CURSO"
                else:
                    horas = float(turnos_dia["DURACAO_HORAS"].fillna(0).sum())
                    valor = "-" if horas == 0 else round(horas, 4)
                    total_horas += horas
            elif data_celula > hoje:
                valor = ""
            elif data_celula in feriados:
                valor = "F"
            elif data_celula.weekday() == 5:
                valor = "S"
            elif data_celula.weekday() == 6:
                valor = "D"
            else:
                valor = "-"
            linha[dia] = valor
        linha["TOTAL (h)"] = "-" if total_horas == 0 else round(total_horas, 4)
        linhas.append(linha)
    return pd.DataFrame(linhas).set_index("EQUIPE") if linhas else pd.DataFrame()


def estilo_horas(valor):
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        minutos_totais = round(float(valor) * 60)
        if minutos_totais >= 8 * 60:
            return "background-color: #16a34a; color: white; font-weight: 700;"
        if minutos_totais < 6 * 60:
            return "background-color: #7e22ce; color: white; font-weight: 700;"
        return "background-color: #dc2626; color: white; font-weight: 700;"
    estilos = {
        "EM CURSO": "background-color: #facc15; color: #713f12; font-weight: 700;",
        "-": "background-color: #ffffff; color: #334155; font-weight: 700;",
        "S": "background-color: #dbeafe; color: #1e3a8a; font-weight: 700;",
        "D": "background-color: #e2e8f0; color: #334155; font-weight: 700;",
        "F": "background-color: #fef3c7; color: #92400e; font-weight: 700;",
        "": "background-color: #f8fafc; color: #94a3b8;",
    }
    return estilos.get(str(valor), "")


def estilo_total_horas(valor):
    if str(valor) == "-":
        return "background-color: #ffffff; color: #334155; font-weight: 700;"
    return "background-color: #f1f5f9; font-weight: 700;"


def montar_tabela_intervalos(dados, ano, mes, grupos, equipes_selecionadas):
    quantidade_dias = calendar.monthrange(ano, mes)[1]
    dias = list(range(1, quantidade_dias + 1))
    hoje = datetime.now(FUSO_GOIAS).date()
    feriados = feriados_brasil(ano)

    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes_selecionadas:
        universo = universo[universo["PREFIXO"].isin(equipes_selecionadas)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    equipes_tabela = sorted(universo["PREFIXO"].dropna().unique())

    datas = pd.to_datetime(universo["DATA"], errors="coerce")
    turnos_mes = universo[datas.dt.year.eq(ano) & datas.dt.month.eq(mes)].copy()
    turnos_mes["DIA"] = pd.to_datetime(turnos_mes["DATA"]).dt.day

    linhas = []
    for equipe in equipes_tabela:
        linha = {"EQUIPE": equipe}
        total_intervalo = 0.0
        turnos_equipe = turnos_mes[turnos_mes["PREFIXO"].eq(equipe)]
        for dia in dias:
            data_celula = date(ano, mes, dia)
            turnos_dia = turnos_equipe[turnos_equipe["DIA"].eq(dia)]
            if not turnos_dia.empty:
                intervalo = float(turnos_dia["INTERVALO_HORAS"].fillna(0).sum())
                valor = round(intervalo, 4)
                total_intervalo += intervalo
            elif data_celula > hoje:
                valor = ""
            elif data_celula in feriados:
                valor = "F"
            elif data_celula.weekday() == 5:
                valor = "S"
            elif data_celula.weekday() == 6:
                valor = "D"
            else:
                valor = "SEM TURNO"
            linha[dia] = valor
        linha["TOTAL"] = round(total_intervalo, 4)
        linhas.append(linha)
    return pd.DataFrame(linhas).set_index("EQUIPE") if linhas else pd.DataFrame()


def estilo_intervalos(valor):
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        minutos_totais = round(float(valor) * 60)
        if minutos_totais > 2 * 60 + 30:
            return "background-color: #7e22ce; color: white; font-weight: 700;"
        if minutos_totais > 1 * 60 + 15:
            return "background-color: #dc2626; color: white; font-weight: 700;"
        if minutos_totais >= 1 * 60:
            return "background-color: #16a34a; color: white; font-weight: 700;"
        return "background-color: #fed7aa; color: #9a3412; font-weight: 700;"
    estilos = {
        "SEM TURNO": "background-color: #ffffff; color: #334155; font-weight: 700;",
        "S": "background-color: #dbeafe; color: #1e3a8a; font-weight: 700;",
        "D": "background-color: #e2e8f0; color: #334155; font-weight: 700;",
        "F": "background-color: #fef3c7; color: #92400e; font-weight: 700;",
        "": "background-color: #f8fafc; color: #94a3b8;",
    }
    return estilos.get(str(valor), "")


def montar_tabela_refeicao(
    dados, intervalos_individuais, ano, mes, grupos, equipes_selecionadas
):
    """Soma somente os intervalos oficiais de refeição por equipe e dia."""
    quantidade_dias = calendar.monthrange(ano, mes)[1]
    dias = list(range(1, quantidade_dias + 1))
    hoje = datetime.now(FUSO_GOIAS).date()
    feriados = feriados_brasil(ano)

    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes_selecionadas:
        universo = universo[universo["PREFIXO"].isin(equipes_selecionadas)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    equipes_tabela = sorted(universo["PREFIXO"].dropna().unique())

    datas_turnos = pd.to_datetime(universo["DATA"], errors="coerce")
    turnos_mes = universo[
        datas_turnos.dt.year.eq(ano) & datas_turnos.dt.month.eq(mes)
    ].copy()
    turnos_mes["DIA"] = pd.to_datetime(turnos_mes["DATA"]).dt.day
    dias_com_turno = set(zip(turnos_mes["PREFIXO"], turnos_mes["DIA"]))

    refeicoes_por_dia = {}
    if not intervalos_individuais.empty:
        refeicoes = intervalos_individuais.copy()
        datas_refeicao = pd.to_datetime(refeicoes["INICIO_INTERVALO"], errors="coerce")
        motivos = refeicoes["MOTIVO_INTERVALO"].fillna("").map(normalizar_nome)
        refeicoes = refeicoes[
            datas_refeicao.dt.year.eq(ano)
            & datas_refeicao.dt.month.eq(mes)
            & refeicoes["GRUPO"].isin(grupos)
            & motivos.eq("REFEICAO")
        ].copy()
        if equipes_selecionadas:
            refeicoes = refeicoes[refeicoes["PREFIXO"].isin(equipes_selecionadas)]
        refeicoes["DIA"] = pd.to_datetime(refeicoes["INICIO_INTERVALO"]).dt.day
        refeicoes_por_dia = (
            refeicoes.groupby(["PREFIXO", "DIA"])["INTERVALO_HORAS"].sum().to_dict()
        )

    linhas = []
    for equipe in equipes_tabela:
        linha = {"EQUIPE": equipe}
        total_refeicao = 0.0
        for dia in dias:
            data_celula = date(ano, mes, dia)
            if (equipe, dia) in dias_com_turno:
                refeicao = float(refeicoes_por_dia.get((equipe, dia), 0.0))
                valor = round(refeicao, 4)
                total_refeicao += refeicao
            elif data_celula > hoje:
                valor = ""
            elif data_celula in feriados:
                valor = "F"
            elif data_celula.weekday() == 5:
                valor = "S"
            elif data_celula.weekday() == 6:
                valor = "D"
            else:
                valor = "SEM TURNO"
            linha[dia] = valor
        linha["TOTAL"] = round(total_refeicao, 4)
        linhas.append(linha)
    return pd.DataFrame(linhas).set_index("EQUIPE") if linhas else pd.DataFrame()


def estilo_refeicao(valor):
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        minutos_totais = round(float(valor) * 60)
        if minutos_totais > 99:
            return "background-color: #7e22ce; color: white; font-weight: 700;"
        if minutos_totais > 75:
            return "background-color: #dc2626; color: white; font-weight: 700;"
        if minutos_totais >= 59:
            return "background-color: #16a34a; color: white; font-weight: 700;"
        return "background-color: #fed7aa; color: #9a3412; font-weight: 700;"
    estilos = {
        "SEM TURNO": "background-color: #ffffff; color: #334155; font-weight: 700;",
        "S": "background-color: #dbeafe; color: #1e3a8a; font-weight: 700;",
        "D": "background-color: #e2e8f0; color: #334155; font-weight: 700;",
        "F": "background-color: #fef3c7; color: #92400e; font-weight: 700;",
        "": "background-color: #f8fafc; color: #94a3b8;",
    }
    return estilos.get(str(valor), "")


def marcar_ds_estimado_gooh(dados):
    """Marca no máximo um possível DS por semana para cada equipe GOOH.

    A fonte ainda não possui uma coluna de escala/DS. Por isso, a estimativa usa a
    abertura mais próxima das 09:00, dentro da janela de 08:30 a 10:30. A regra é
    aplicada sobre todo o histórico carregado para não duplicar um DS na virada do mês.
    """
    resultado = pd.Series(False, index=dados.index, dtype=bool)
    if dados.empty:
        return resultado

    inicio = pd.to_datetime(dados["INICIO_TURNO"], errors="coerce")
    minutos_abertura = inicio.dt.hour * 60 + inicio.dt.minute
    candidatos = dados[
        dados["GRUPO"].eq("GOOH")
        & inicio.notna()
        & minutos_abertura.between(8 * 60 + 30, 10 * 60 + 30)
    ].copy()
    if candidatos.empty:
        return resultado

    inicio_candidatos = pd.to_datetime(candidatos["INICIO_TURNO"], errors="coerce")
    calendario_iso = inicio_candidatos.dt.isocalendar()
    candidatos["ANO_ISO"] = calendario_iso.year.astype(int)
    candidatos["SEMANA_ISO"] = calendario_iso.week.astype(int)
    candidatos["DISTANCIA_09H"] = (
        inicio_candidatos.dt.hour * 60 + inicio_candidatos.dt.minute - 9 * 60
    ).abs()
    escolhidos = (
        candidatos.sort_values(["DISTANCIA_09H", "INICIO_TURNO"])
        .groupby(["PREFIXO", "ANO_ISO", "SEMANA_ISO"], sort=False)
        .head(1)
        .index
    )
    resultado.loc[escolhidos] = True
    return resultado


def calcular_analise_mensal(dados, ano, mes, grupos, equipes_selecionadas):
    """Calcula o percentual de jornadas abaixo da meta para cada equipe."""
    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes_selecionadas:
        universo = universo[universo["PREFIXO"].isin(equipes_selecionadas)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    universo["DS_ESTIMADO"] = marcar_ds_estimado_gooh(dados).reindex(
        universo.index, fill_value=False
    )

    datas = pd.to_datetime(universo["DATA"], errors="coerce")
    jornadas = universo[
        datas.dt.year.eq(ano)
        & datas.dt.month.eq(mes)
        & universo["FIM_TURNO"].notna()
        & universo["DURACAO_HORAS"].notna()
    ].copy()
    if jornadas.empty:
        return pd.DataFrame(), jornadas

    jornadas["META_HORAS"] = META_HORAS_PADRAO
    jornadas.loc[jornadas["GRUPO"].eq("GOOH"), "META_HORAS"] = META_HORAS_GOOH
    jornadas.loc[
        jornadas["GRUPO"].eq("GOOH") & jornadas["DS_ESTIMADO"], "META_HORAS"
    ] = META_HORAS_GOOH_DS
    jornadas["DEFICIT_HORAS"] = (
        jornadas["META_HORAS"] - jornadas["DURACAO_HORAS"]
    ).clip(lower=0)
    jornadas["COM_DESVIO"] = jornadas["DEFICIT_HORAS"].gt(1 / 60)

    ranking = (
        jornadas.groupby(["GRUPO", "PREFIXO"], as_index=False)
        .agg(
            JORNADAS_ANALISADAS=("PREFIXO", "size"),
            JORNADAS_COM_DESVIO=("COM_DESVIO", "sum"),
            MEDIA_HORAS=("DURACAO_HORAS", "mean"),
            DEFICIT_TOTAL_HORAS=("DEFICIT_HORAS", "sum"),
            DIAS_DS_ESTIMADOS=("DS_ESTIMADO", "sum"),
        )
    )
    ranking["PERCENTUAL_DESVIO"] = (
        100 * ranking["JORNADAS_COM_DESVIO"] / ranking["JORNADAS_ANALISADAS"]
    ).round(1)
    ranking["MEDIA_HORAS"] = ranking["MEDIA_HORAS"].round(4)
    ranking["DEFICIT_TOTAL_HORAS"] = ranking["DEFICIT_TOTAL_HORAS"].round(4)
    ranking = ranking.sort_values(
        ["GRUPO", "PERCENTUAL_DESVIO", "DEFICIT_TOTAL_HORAS", "PREFIXO"],
        ascending=[True, False, False, True],
    ).reset_index(drop=True)
    return ranking, jornadas



def preparar_relatorio_equipe(dados, intervalos_individuais, equipe, ano, mes):
    """Reúne todos os turnos e intervalos do mês para uma única equipe."""
    datas = pd.to_datetime(dados["DATA"], errors="coerce")
    turnos = dados[
        dados["PREFIXO"].eq(equipe)
        & datas.dt.year.eq(ano)
        & datas.dt.month.eq(mes)
    ].copy()
    turnos["DS_ESTIMADO"] = marcar_ds_estimado_gooh(dados).reindex(
        turnos.index, fill_value=False
    )
    turnos["META_HORAS"] = META_HORAS_PADRAO
    turnos.loc[turnos["GRUPO"].eq("GOOH"), "META_HORAS"] = META_HORAS_GOOH
    turnos.loc[
        turnos["GRUPO"].eq("GOOH") & turnos["DS_ESTIMADO"], "META_HORAS"
    ] = META_HORAS_GOOH_DS
    turnos["RESULTADO"] = "EM CURSO"
    encerrados = turnos["FIM_TURNO"].notna() & turnos["DURACAO_HORAS"].notna()
    turnos.loc[
        encerrados & turnos["DURACAO_HORAS"].ge(turnos["META_HORAS"]), "RESULTADO"
    ] = "DENTRO DA META"
    turnos.loc[
        encerrados & turnos["DURACAO_HORAS"].lt(turnos["META_HORAS"]), "RESULTADO"
    ] = "ABAIXO DA META"
    turnos["REGRA_DIA"] = "Meta padrão"
    turnos.loc[turnos["DS_ESTIMADO"], "REGRA_DIA"] = "DS estimado"

    intervalos = intervalos_individuais[
        intervalos_individuais["PREFIXO"].eq(equipe)
    ].copy() if not intervalos_individuais.empty else pd.DataFrame()
    if not intervalos.empty:
        inicio_intervalo = pd.to_datetime(intervalos["INICIO_INTERVALO"], errors="coerce")
        intervalos = intervalos[
            inicio_intervalo.dt.year.eq(ano) & inicio_intervalo.dt.month.eq(mes)
        ].copy()
        intervalos["DATA_INTERVALO"] = pd.to_datetime(
            intervalos["INICIO_INTERVALO"], errors="coerce"
        ).dt.date
        intervalos["CATEGORIA"] = intervalos["MOTIVO_INTERVALO"].fillna("").map(
            normalizar_nome
        )

    colunas_intervalo = [
        "NUM_INTERVALOS", "TOTAL_INTERVALOS_HORAS", "REFEICAO_HORAS",
        "MANUTENCAO_HORAS", "RETORNO_BASE_HORAS",
    ]
    for coluna in colunas_intervalo:
        turnos[coluna] = 0.0

    if not intervalos.empty:
        totais = intervalos.groupby("DATA_INTERVALO").agg(
            NUM_INTERVALOS=("INTERVALO_ID", "nunique"),
            TOTAL_INTERVALOS_HORAS=("INTERVALO_HORAS", "sum"),
        )
        refeicao = intervalos[intervalos["CATEGORIA"].eq("REFEICAO")].groupby(
            "DATA_INTERVALO"
        )["INTERVALO_HORAS"].sum()
        manutencao = intervalos[
            intervalos["CATEGORIA"].str.contains("MANUTENCAO", na=False)
        ].groupby("DATA_INTERVALO")["INTERVALO_HORAS"].sum()
        retorno = intervalos[
            intervalos["CATEGORIA"].str.contains("RETORNO.*BASE", regex=True, na=False)
        ].groupby("DATA_INTERVALO")["INTERVALO_HORAS"].sum()
        totais["REFEICAO_HORAS"] = refeicao
        totais["MANUTENCAO_HORAS"] = manutencao
        totais["RETORNO_BASE_HORAS"] = retorno
        totais = totais.fillna(0)
        for indice in turnos.index:
            data_turno = turnos.at[indice, "DATA"]
            if data_turno in totais.index:
                for coluna in colunas_intervalo:
                    turnos.at[indice, coluna] = float(totais.at[data_turno, coluna])

    turnos["DIFERENCA_SAIDA_HORAS"] = pd.to_numeric(
        turnos["DIFERENCA_FECHAMENTO_MIN"], errors="coerce"
    ) / 60
    turnos["PERCENTUAL_INTERVALO"] = 0.0
    duracao_valida = turnos["DURACAO_HORAS"].fillna(0).gt(0)
    turnos.loc[duracao_valida, "PERCENTUAL_INTERVALO"] = (
        turnos.loc[duracao_valida, "TOTAL_INTERVALOS_HORAS"]
        / turnos.loc[duracao_valida, "DURACAO_HORAS"]
    )

    def situacao_almoco(horas):
        minutos = round(float(horas or 0) * 60)
        if minutos == 0:
            return "Sem refeição registrada"
        if minutos < 59:
            return "Abaixo de 00:59"
        if minutos <= 75:
            return "Dentro de 00:59 a 01:15"
        return "Acima de 01:15"

    turnos["SITUACAO_ALMOCO"] = turnos["REFEICAO_HORAS"].map(situacao_almoco)
    return turnos.sort_values("INICIO_TURNO"), intervalos.sort_values(
        "INICIO_INTERVALO"
    ) if not intervalos.empty else intervalos


@st.cache_data(show_spinner=False)
def gerar_relatorio_equipe_excel(dados, intervalos_individuais, equipe, ano, mes):
    """Gera um Excel individual com todos os horários da equipe no mês."""
    turnos, intervalos = preparar_relatorio_equipe(
        dados, intervalos_individuais, equipe, ano, mes
    )
    livro = Workbook()
    planilha = livro.active
    planilha.title = "Turnos do mês"

    azul_escuro = "1F4E78"
    azul_cabecalho = "2E75B6"
    azul_claro = "D9EAF7"
    branco = "FFFFFF"
    verde = "C6EFCE"
    vermelho = "FFC7CE"
    amarelo = "FFF2CC"
    cinza = "E7E6E6"
    borda_fina = Side(style="thin", color="B7C9D6")

    planilha.merge_cells("A1:S1")
    planilha["A1"] = f"Análise de turnos e intervalos — {equipe}"
    planilha["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
    planilha["A1"].font = Font(color=branco, bold=True, size=16)
    planilha["A1"].alignment = Alignment(vertical="center")
    planilha.row_dimensions[1].height = 30
    planilha.merge_cells("A2:S2")
    planilha["A2"] = (
        f"Período: {MESES[mes - 1]} de {ano}. Jornada = primeira abertura até o "
        "último fechamento do dia; os intervalos não são descontados da duração."
    )
    planilha["A2"].fill = PatternFill("solid", fgColor=azul_claro)
    planilha["A2"].font = Font(color=azul_escuro, italic=True)

    encerrados = turnos[turnos["FIM_TURNO"].notna()]
    desvios = int(encerrados["RESULTADO"].eq("ABAIXO DA META").sum())
    taxa_desvio = desvios / len(encerrados) if len(encerrados) else 0
    total_intervalos = int(turnos["NUM_INTERVALOS"].sum()) if not turnos.empty else 0
    horas_intervalo = float(turnos["TOTAL_INTERVALOS_HORAS"].sum()) if not turnos.empty else 0
    horas_turno = float(turnos["DURACAO_HORAS"].fillna(0).sum()) if not turnos.empty else 0
    percentual_intervalos = horas_intervalo / horas_turno if horas_turno else 0
    metricas = [
        ("Equipe", equipe),
        ("Turnos", len(turnos)),
        ("Intervalos", total_intervalos),
        ("Total intervalos", horas_intervalo / 24),
        ("Intervalos / turno", percentual_intervalos),
        ("Jornadas com desvio", taxa_desvio),
    ]
    posicoes = [("A4", "B4"), ("D4", "E4"), ("G4", "H4"),
                ("J4", "K4"), ("M4", "N4"), ("P4", "Q4")]
    for (rotulo, valor), (celula_rotulo, celula_valor) in zip(metricas, posicoes):
        planilha[celula_rotulo] = rotulo
        planilha[celula_valor] = valor
        planilha[celula_rotulo].fill = PatternFill("solid", fgColor=azul_claro)
        planilha[celula_rotulo].font = Font(color=azul_escuro, bold=True)
        planilha[celula_valor].font = Font(bold=True)
        for celula in (planilha[celula_rotulo], planilha[celula_valor]):
            celula.border = Border(
                left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina
            )
            celula.alignment = Alignment(vertical="center")
    planilha["K4"].number_format = "[h]:mm"
    planilha["N4"].number_format = "0.0%"
    planilha["Q4"].number_format = "0.0%"

    cabecalhos = [
        "Equipe", "Data", "Abertura", "Saída prevista", "Fechamento",
        "Duração do turno", "Diferença da saída", "Nº aberturas", "Nº intervalos",
        "Total de intervalos", "% do turno", "Refeição", "Manutenção",
        "Retorno à base", "Situação do almoço", "Meta do dia", "Regra do dia",
        "Resultado", "Observação",
    ]
    linha_cabecalho = 7
    for coluna, cabecalho in enumerate(cabecalhos, 1):
        celula = planilha.cell(linha_cabecalho, coluna, cabecalho)
        celula.fill = PatternFill("solid", fgColor=azul_cabecalho)
        celula.font = Font(color=branco, bold=True)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = Border(
            left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina
        )
    planilha.row_dimensions[linha_cabecalho].height = 34

    for numero_linha, (_, item) in enumerate(turnos.iterrows(), linha_cabecalho + 1):
        valores = [
            item["PREFIXO"], item["DATA"], item["INICIO_TURNO"], item["SAIDA_PREVISTA"],
            item["FIM_TURNO"], item["DURACAO_HORAS"] / 24 if pd.notna(item["DURACAO_HORAS"]) else None,
            formatar_duracao(item["DIFERENCA_SAIDA_HORAS"])
            if pd.notna(item["DIFERENCA_SAIDA_HORAS"]) else "",
            int(item["ABERTURAS_NO_DIA"]), int(item["NUM_INTERVALOS"]),
            item["TOTAL_INTERVALOS_HORAS"] / 24, item["PERCENTUAL_INTERVALO"],
            item["REFEICAO_HORAS"] / 24, item["MANUTENCAO_HORAS"] / 24,
            item["RETORNO_BASE_HORAS"] / 24, item["SITUACAO_ALMOCO"],
            item["META_HORAS"] / 24, item["REGRA_DIA"], item["RESULTADO"],
            item.get("OBSERVACAO", ""),
        ]
        for coluna, valor in enumerate(valores, 1):
            if pd.isna(valor):
                valor = None
            celula = planilha.cell(numero_linha, coluna, valor)
            celula.border = Border(
                left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina
            )
            celula.alignment = Alignment(vertical="center", wrap_text=True)
        for coluna_data in (2,):
            planilha.cell(numero_linha, coluna_data).number_format = "dd/mm/yyyy"
        for coluna_data_hora in (3, 4, 5):
            planilha.cell(numero_linha, coluna_data_hora).number_format = "dd/mm/yyyy hh:mm"
        for coluna_duracao in (6, 10, 12, 13, 14, 16):
            planilha.cell(numero_linha, coluna_duracao).number_format = "[h]:mm;-[h]:mm"
        planilha.cell(numero_linha, 11).number_format = "0.0%"
        resultado = planilha.cell(numero_linha, 18)
        if item["RESULTADO"] == "ABAIXO DA META":
            resultado.fill = PatternFill("solid", fgColor=vermelho)
        elif item["RESULTADO"] == "DENTRO DA META":
            resultado.fill = PatternFill("solid", fgColor=verde)
        else:
            resultado.fill = PatternFill("solid", fgColor=amarelo)

    ultima_linha = max(linha_cabecalho + 1, linha_cabecalho + len(turnos))
    if not turnos.empty:
        tabela = Table(displayName="TabelaTurnos", ref=f"A7:S{ultima_linha}")
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        planilha.add_table(tabela)
    planilha.freeze_panes = "A8"
    planilha.auto_filter.ref = f"A7:S{ultima_linha}"
    larguras = [16, 12, 19, 19, 19, 18, 18, 14, 14, 19, 13, 13, 14, 16, 25, 14, 16, 19, 35]
    for indice, largura in enumerate(larguras, 1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    planilha_intervalos = livro.create_sheet("Intervalos individuais")
    planilha_intervalos.merge_cells("A1:F1")
    planilha_intervalos["A1"] = f"Intervalos individuais — {equipe}"
    planilha_intervalos["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
    planilha_intervalos["A1"].font = Font(color=branco, bold=True, size=15)
    cabecalhos_intervalos = ["Data", "Início", "Fim", "Duração", "Motivo", "ID do intervalo"]
    for coluna, cabecalho in enumerate(cabecalhos_intervalos, 1):
        celula = planilha_intervalos.cell(3, coluna, cabecalho)
        celula.fill = PatternFill("solid", fgColor=azul_cabecalho)
        celula.font = Font(color=branco, bold=True)
        celula.alignment = Alignment(horizontal="center")
    if not intervalos.empty:
        for numero_linha, (_, item) in enumerate(intervalos.iterrows(), 4):
            valores = [
                item["DATA_INTERVALO"], item["INICIO_INTERVALO"], item["FIM_INTERVALO"],
                item["INTERVALO_HORAS"] / 24, item["MOTIVO_INTERVALO"], item["INTERVALO_ID"],
            ]
            for coluna, valor in enumerate(valores, 1):
                planilha_intervalos.cell(numero_linha, coluna, None if pd.isna(valor) else valor)
            planilha_intervalos.cell(numero_linha, 1).number_format = "dd/mm/yyyy"
            planilha_intervalos.cell(numero_linha, 2).number_format = "dd/mm/yyyy hh:mm"
            planilha_intervalos.cell(numero_linha, 3).number_format = "dd/mm/yyyy hh:mm"
            planilha_intervalos.cell(numero_linha, 4).number_format = "[h]:mm"
        ultima_linha_intervalos = 3 + len(intervalos)
        tabela_intervalos = Table(
            displayName="TabelaIntervalos", ref=f"A3:F{ultima_linha_intervalos}"
        )
        tabela_intervalos.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        planilha_intervalos.add_table(tabela_intervalos)
    planilha_intervalos.freeze_panes = "A4"
    for indice, largura in enumerate([13, 21, 21, 14, 30, 18], 1):
        planilha_intervalos.column_dimensions[get_column_letter(indice)].width = largura

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def montar_painel_tempos(dados, intervalos_individuais, ano, mes, grupos, equipes):
    """Compõe trabalho líquido e intervalos por equipe/mês e por equipe/dia."""
    universo = dados[dados["GRUPO"].isin(grupos)].copy()
    if equipes:
        universo = universo[universo["PREFIXO"].isin(equipes)]
    universo = ocultar_desmobilizadas_sem_movimento(universo, ano, mes)
    equipes_ativas = universo[["GRUPO", "PREFIXO"]].drop_duplicates()

    datas_turnos = pd.to_datetime(universo["DATA"], errors="coerce")
    turnos_mes = universo[
        datas_turnos.dt.year.eq(ano) & datas_turnos.dt.month.eq(mes)
    ].copy()
    colunas_tempos = [
        "REFEICAO_HORAS", "MANUTENCAO_HORAS", "RETORNO_BASE_HORAS",
        "OUTROS_INTERVALOS_HORAS", "TOTAL_INTERVALOS_HORAS",
    ]
    for coluna in colunas_tempos:
        turnos_mes[coluna] = 0.0

    intervalos_mes = pd.DataFrame()
    if not intervalos_individuais.empty:
        intervalos_mes = intervalos_individuais[
            intervalos_individuais["GRUPO"].isin(grupos)
        ].copy()
        if equipes:
            intervalos_mes = intervalos_mes[
                intervalos_mes["PREFIXO"].isin(equipes)
            ]
        inicio_intervalo = pd.to_datetime(
            intervalos_mes["INICIO_INTERVALO"], errors="coerce"
        )
        intervalos_mes = intervalos_mes[
            inicio_intervalo.dt.year.eq(ano) & inicio_intervalo.dt.month.eq(mes)
        ].copy()
        if not intervalos_mes.empty:
            intervalos_mes["DATA"] = pd.to_datetime(
                intervalos_mes["INICIO_INTERVALO"], errors="coerce"
            ).dt.date
            intervalos_mes["CATEGORIA"] = intervalos_mes[
                "MOTIVO_INTERVALO"
            ].fillna("").map(normalizar_nome)
            intervalos_mes["TIPO_PAINEL"] = "OUTROS"
            intervalos_mes.loc[
                intervalos_mes["CATEGORIA"].eq("REFEICAO"), "TIPO_PAINEL"
            ] = "REFEICAO"
            intervalos_mes.loc[
                intervalos_mes["CATEGORIA"].str.contains("MANUTENCAO", na=False),
                "TIPO_PAINEL",
            ] = "MANUTENCAO"
            intervalos_mes.loc[
                intervalos_mes["CATEGORIA"].str.contains(
                    "RETORNO.*BASE", regex=True, na=False
                ),
                "TIPO_PAINEL",
            ] = "RETORNO_BASE"
            por_tipo = intervalos_mes.pivot_table(
                index=["PREFIXO", "DATA"], columns="TIPO_PAINEL",
                values="INTERVALO_HORAS", aggfunc="sum", fill_value=0,
            )
            por_tipo = por_tipo.rename(columns={
                "REFEICAO": "REFEICAO_HORAS",
                "MANUTENCAO": "MANUTENCAO_HORAS",
                "RETORNO_BASE": "RETORNO_BASE_HORAS",
                "OUTROS": "OUTROS_INTERVALOS_HORAS",
            }).reset_index()
            turnos_mes = turnos_mes.merge(
                por_tipo, on=["PREFIXO", "DATA"], how="left", suffixes=("", "_CALC")
            )
            for coluna in colunas_tempos[:-1]:
                calculada = f"{coluna}_CALC"
                if calculada in turnos_mes.columns:
                    turnos_mes[coluna] = turnos_mes[calculada].fillna(0)
                    turnos_mes = turnos_mes.drop(columns=calculada)

    categorias = colunas_tempos[:-1]
    total_categorizado = turnos_mes[categorias].sum(axis=1)
    if "INTERVALO_HORAS" in turnos_mes.columns:
        total_oficial = pd.to_numeric(
            turnos_mes["INTERVALO_HORAS"], errors="coerce"
        ).fillna(0)
        residual = (total_oficial - total_categorizado).clip(lower=0)
        turnos_mes["OUTROS_INTERVALOS_HORAS"] += residual
        turnos_mes["TOTAL_INTERVALOS_HORAS"] = total_oficial.combine(
            turnos_mes[categorias].sum(axis=1), max
        )
    else:
        turnos_mes["TOTAL_INTERVALOS_HORAS"] = total_categorizado
    # Tempo em atividade é usado somente neste painel de composição. A aba de horas
    # trabalhadas continua mostrando a jornada completa, sem desconto de intervalos.
    turnos_mes["TRABALHO_HORAS"] = (
        turnos_mes["DURACAO_HORAS"].fillna(0)
        - turnos_mes["TOTAL_INTERVALOS_HORAS"]
    ).clip(lower=0)

    agregacoes = {
        "TRABALHO_HORAS": "sum",
        "REFEICAO_HORAS": "sum",
        "MANUTENCAO_HORAS": "sum",
        "RETORNO_BASE_HORAS": "sum",
        "OUTROS_INTERVALOS_HORAS": "sum",
        "TOTAL_INTERVALOS_HORAS": "sum",
    }
    if turnos_mes.empty:
        resumo = equipes_ativas.copy()
        for coluna in agregacoes:
            resumo[coluna] = 0.0
    else:
        resumo_calculado = turnos_mes.groupby(
            ["GRUPO", "PREFIXO"], as_index=False
        ).agg(agregacoes)
        resumo = equipes_ativas.merge(
            resumo_calculado, on=["GRUPO", "PREFIXO"], how="left"
        )
        resumo[list(agregacoes)] = resumo[list(agregacoes)].fillna(0)

    resumo["PRIORIDADE"] = 0
    resumo["ALERTAS"] = ""
    if not intervalos_mes.empty:
        for indice, equipe_linha in resumo.iterrows():
            intervalos_equipe = intervalos_mes[
                intervalos_mes["PREFIXO"].eq(equipe_linha["PREFIXO"])
            ]
            manutencoes = intervalos_equipe[
                intervalos_equipe["TIPO_PAINEL"].eq("MANUTENCAO")
            ]
            refeicoes = intervalos_equipe[
                intervalos_equipe["TIPO_PAINEL"].eq("REFEICAO")
            ]
            retornos = intervalos_equipe[
                intervalos_equipe["TIPO_PAINEL"].eq("RETORNO_BASE")
            ]
            manutencao_90 = int(manutencoes["INTERVALO_HORAS"].gt(1.5).sum())
            manutencao_60 = int(manutencoes["INTERVALO_HORAS"].ge(1.0).sum())
            refeicao_75 = int(refeicoes["INTERVALO_HORAS"].gt(1.25).sum())
            retorno_90 = int(retornos["INTERVALO_HORAS"].gt(1.5).sum())
            alertas = []
            prioridade = 0
            if manutencao_90:
                prioridade = 1
                alertas.append(f"{manutencao_90} manutenção(ões) acima de 01:30")
            if manutencao_60 > 1:
                prioridade = 1
                alertas.append(f"{manutencao_60} manutenções de 01:00 ou mais")
            if refeicao_75 > 2:
                prioridade = prioridade or 2
                alertas.append(f"{refeicao_75} refeições acima de 01:15")
            if retorno_90:
                prioridade = prioridade or 3
                alertas.append(f"{retorno_90} retorno(s) à base acima de 01:30")
            resumo.at[indice, "PRIORIDADE"] = prioridade
            resumo.at[indice, "ALERTAS"] = " · ".join(alertas)

    resumo = resumo.sort_values(
        ["GRUPO", "TOTAL_INTERVALOS_HORAS", "PREFIXO"],
        ascending=[True, False, True],
    ).reset_index(drop=True)
    return resumo, turnos_mes.sort_values(["PREFIXO", "DATA"]), intervalos_mes


def html_barra_tempos(trabalho, refeicao, manutencao, retorno, outros=0):
    """Renderiza uma barra horizontal empilhada com valores auditáveis."""
    componentes = [
        ("Trabalhando", float(trabalho or 0), "#16a34a", "#ffffff"),
        ("Refeição", float(refeicao or 0), "#2563eb", "#ffffff"),
        ("Manutenção", float(manutencao or 0), "#dc2626", "#ffffff"),
        ("Retorno à base", float(retorno or 0), "#eab308", "#422006"),
        ("Outros intervalos", float(outros or 0), "#94a3b8", "#0f172a"),
    ]
    total = sum(valor for _, valor, _, _ in componentes)
    if total <= 0:
        return (
            '<div class="tempo-barra tempo-vazia"></div>'
            '<div class="tempo-valores">Sem turno registrado</div>'
        )
    segmentos = []
    valores = []
    for nome, valor, cor, cor_texto in componentes:
        if valor <= 0:
            continue
        largura = 100 * valor / total
        segmentos.append(
            f'<div class="tempo-segmento" style="width:{largura:.4f}%;'
            f'background:{cor};color:{cor_texto}" title="{nome}: '
            f'{formatar_duracao(valor)}"></div>'
        )
        valores.append(
            f'<span><i style="background:{cor}"></i>{nome} '
            f'<strong>{formatar_duracao(valor)}</strong></span>'
        )
    return (
        '<div class="tempo-barra">' + "".join(segmentos) + '</div>'
        '<div class="tempo-valores">' + "".join(valores) + '</div>'
    )


st.title("Acompanhamento de Turnos GO")
st.caption("Abertura e fechamento reais • dados atualizados pelo bot")
if st.sidebar.button("Atualizar dados", type="primary", use_container_width=True):
    carregar_turnos.clear()
    st.rerun()

try:
    with st.spinner("Lendo a planilha de turnos..."):
        brutos, atualizado_em = carregar_turnos()
        if brutos.empty:
            st.warning("A planilha ainda está vazia. Execute o teste do bot para enviar os dados.")
            st.stop()
        dados = preparar_dados(brutos)
        intervalos_individuais = preparar_intervalos_individuais(brutos)
except Exception as erro:
    st.error("Não foi possível consultar a planilha de turnos.")
    st.code(str(erro))
    st.stop()

if atualizado_em:
    try:
        atualizado = datetime.fromisoformat(atualizado_em).astimezone(FUSO_GOIAS)
        st.caption(f"Última carga do bot: {atualizado:%d/%m/%Y %H:%M:%S}")
    except ValueError:
        st.caption("Última carga do bot: " + atualizado_em)

datas_validas = pd.to_datetime(dados["DATA"], errors="coerce").dropna()
ano_padrao = int(datas_validas.dt.year.max()) if not datas_validas.empty else date.today().year
meses = datas_validas[datas_validas.dt.year.eq(ano_padrao)].dt.month
mes_padrao = int(meses.max()) if not meses.empty else date.today().month
ano = st.sidebar.number_input("Ano", min_value=2020, max_value=2100, value=ano_padrao)
mes = st.sidebar.selectbox(
    "Mês", range(1, 13), index=mes_padrao - 1,
    format_func=lambda valor: MESES[valor - 1],
)

datas = pd.to_datetime(dados["DATA"], errors="coerce")
filtrado = dados[datas.dt.year.eq(ano) & datas.dt.month.eq(mes)].copy()
grupos = st.sidebar.multiselect("Grupos", GRUPOS, default=GRUPOS)
filtrado = filtrado[filtrado["GRUPO"].isin(grupos)]
equipes = st.sidebar.multiselect("Equipes", sorted(filtrado["PREFIXO"].dropna().unique()))
if equipes:
    filtrado = filtrado[filtrado["PREFIXO"].isin(equipes)]
situacoes = st.sidebar.multiselect(
    "Situação", ["EM ANDAMENTO", "ENCERRADO"], default=["EM ANDAMENTO", "ENCERRADO"]
)
filtrado = filtrado[filtrado["SITUACAO"].isin(situacoes)]

if intervalos_individuais.empty:
    ranking_intervalos = pd.DataFrame()
else:
    datas_ranking = pd.to_datetime(
        intervalos_individuais["INICIO_INTERVALO"], errors="coerce"
    )
    ranking_intervalos = intervalos_individuais[
        datas_ranking.dt.year.eq(ano)
        & datas_ranking.dt.month.eq(mes)
        & intervalos_individuais["GRUPO"].isin(grupos)
        & intervalos_individuais["INTERVALO_OFICIAL_SEGUNDOS"].gt(0)
    ].copy()
    if equipes:
        ranking_intervalos = ranking_intervalos[
            ranking_intervalos["PREFIXO"].isin(equipes)
        ]
    ranking_intervalos = ranking_intervalos.sort_values(
        ["INTERVALO_OFICIAL_SEGUNDOS", "INICIO_INTERVALO"],
        ascending=[False, True],
    ).reset_index(drop=True)
    ranking_intervalos["MOTIVO_INTERVALO"] = (
        ranking_intervalos["MOTIVO_INTERVALO"].fillna("").astype(str).str.strip()
        .replace("", "Não informado")
    )
    ranking_intervalos["CATEGORIA_MOTIVO"] = ranking_intervalos[
        "MOTIVO_INTERVALO"
    ].map(normalizar_nome)

ranking_mensal, jornadas_mensais = calcular_analise_mensal(
    dados, ano, mes, grupos, equipes
)

metricas = st.columns(4)
metricas[0].metric("Turnos consolidados", filtrado["HIST_TURMA_PLANTAO_ID"].nunique())
metricas[1].metric("Equipes", filtrado["PREFIXO"].nunique())
metricas[2].metric("Em andamento", int(filtrado["SITUACAO"].eq("EM ANDAMENTO").sum()))
mediana = filtrado["DURACAO_HORAS"].median()
metricas[3].metric("Duração mediana", "—" if pd.isna(mediana) else formatar_duracao(mediana))

colunas = [
    "DATA", "GRUPO", "PREFIXO", "INICIO_TURNO", "SAIDA_PREVISTA", "FIM_TURNO",
    "SITUACAO", "ABERTURAS_NO_DIA", "PERMANENCIA", "INTERVALO", "DURACAO",
    "DIFERENCA_FECHAMENTO_MIN",
    "MOTIVOS_INTERVALO", "PARTICIPA_ESCALA", "OBSERVACAO",
]
(
    aba_analise, aba_turnos, aba_mapa, aba_horas, aba_intervalos,
    aba_refeicao, aba_ranking, aba_resumo,
) = st.tabs(
    [
        "Análise mensal", "Turnos", "Mapa mensal", "Horas trabalhadas",
        "Intervalos", "Intervalos de refeição", "Ranking de intervalos",
        "Resumo diário",
    ]
)
with aba_analise:
    st.subheader(f"Análise mensal — {MESES[mes - 1]} de {ano}")
    st.caption(
        "Percentual de desvio = jornadas encerradas abaixo da meta ÷ jornadas "
        "encerradas analisadas. Metas: 08:00 para GOOL, GOOC e GOOK; 09:00 para "
        "GOOH; 08:00 no DS estimado da GOOH."
    )
    with st.expander("O que significa jornada e como o desvio é calculado?"):
        st.markdown(
            """
            **Jornada** é o tempo entre a primeira abertura e o último fechamento da
            equipe no mesmo dia. Quando a equipe reabre o turno, todas as aberturas são
            consolidadas em uma única jornada. Os intervalos aparecem separadamente e
            **não são descontados** da duração da jornada.

            Uma jornada encerrada é marcada com desvio quando dura menos de **08:00**
            para GOOL, GOOC e GOOK, ou menos de **09:00** para GOOH. No possível dia de
            DS da GOOH, a meta considerada é **08:00**. Turnos em andamento, dias futuros
            e dias sem abertura não entram no percentual desta tela.

            Exemplo: 3 jornadas abaixo da meta em 4 jornadas encerradas resultam em
            **75% de desvio**.
            """
        )
    if jornadas_mensais.empty:
        st.info("Ainda não há jornadas encerradas para analisar neste período.")
    else:
        total_jornadas = len(jornadas_mensais)
        total_desvios = int(jornadas_mensais["COM_DESVIO"].sum())
        taxa_geral = 100 * total_desvios / total_jornadas
        deficit_total = float(jornadas_mensais["DEFICIT_HORAS"].sum())
        indicadores = st.columns(4)
        indicadores[0].metric("Equipes analisadas", ranking_mensal["PREFIXO"].nunique())
        indicadores[1].metric("Jornadas encerradas", total_jornadas)
        indicadores[2].metric(
            "Jornadas com desvio", f"{total_desvios} ({taxa_geral:.1f}%)"
        )
        indicadores[3].metric("Déficit acumulado", formatar_duracao(deficit_total))

        st.markdown("#### Equipes mais ofensoras por prefixo")
        for grupo in grupos:
            ranking_grupo = ranking_mensal[
                ranking_mensal["GRUPO"].eq(grupo)
            ].copy()
            if ranking_grupo.empty:
                continue
            with st.container(border=True):
                st.markdown(f"### {grupo}")
                for _, item in ranking_grupo.head(5).iterrows():
                    (
                        coluna_equipe, coluna_barra, coluna_percentual,
                        coluna_relatorio,
                    ) = st.columns(
                        [2.1, 5.2, 1.1, 1.8], vertical_alignment="center"
                    )
                    coluna_equipe.markdown(f"**{item['PREFIXO']}**")
                    percentual = float(item["PERCENTUAL_DESVIO"])
                    coluna_barra.progress(
                        min(100, max(0, round(percentual))),
                        text=(
                            f"{int(item['JORNADAS_COM_DESVIO'])} de "
                            f"{int(item['JORNADAS_ANALISADAS'])} jornadas"
                        ),
                    )
                    coluna_percentual.markdown(f"**{percentual:.1f}%**")
                    arquivo_equipe = gerar_relatorio_equipe_excel(
                        dados, intervalos_individuais, item["PREFIXO"], ano, mes
                    )
                    coluna_relatorio.download_button(
                        "Baixar relatório",
                        arquivo_equipe,
                        file_name=(
                            f"relatorio_{item['PREFIXO']}_{ano}_{mes:02}.xlsx"
                        ),
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"
                        ),
                        key=f"relatorio_{item['PREFIXO']}_{ano}_{mes:02}",
                        use_container_width=True,
                    )

                with st.expander(f"Ver ranking completo de {grupo}"):
                    exibir = ranking_grupo.copy()
                    exibir["MÉDIA"] = exibir["MEDIA_HORAS"].map(formatar_duracao)
                    exibir["DÉFICIT TOTAL"] = exibir["DEFICIT_TOTAL_HORAS"].map(
                        formatar_duracao
                    )
                    exibir["% DESVIO"] = exibir["PERCENTUAL_DESVIO"].map(
                        lambda valor: f"{valor:.1f}%"
                    )
                    st.dataframe(
                        exibir[[
                            "PREFIXO", "% DESVIO", "JORNADAS_COM_DESVIO",
                            "JORNADAS_ANALISADAS", "MÉDIA", "DÉFICIT TOTAL",
                            "DIAS_DS_ESTIMADOS",
                        ]],
                        hide_index=True,
                        use_container_width=True,
                        column_config={
                            "PREFIXO": "Equipe",
                            "JORNADAS_COM_DESVIO": "Jornadas com desvio",
                            "JORNADAS_ANALISADAS": "Jornadas analisadas",
                            "DIAS_DS_ESTIMADOS": "DS estimados",
                        },
                    )

        st.info(
            "DS estimado: no máximo uma abertura GOOH por equipe/semana entre "
            "08:30 e 10:30, escolhendo a mais próxima das 09:00. Quando a fonte "
            "passar a informar o DS, esta estimativa poderá ser substituída pelo dado oficial."
        )
        st.download_button(
            "Baixar ranking mensal em CSV",
            ranking_mensal.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
            file_name=f"ranking_desvios_{ano}_{mes:02}.csv",
            mime="text/csv",
        )
with aba_turnos:
    st.subheader(f"Composição dos turnos — {MESES[mes - 1]} de {ano}")
    st.caption(
        "Todas as equipes ativas aparecem dentro do respectivo prefixo, ordenadas "
        "do maior para o menor total de intervalos. Neste painel, “trabalhando” é a "
        "duração da jornada menos os intervalos oficiais; essa composição não altera "
        "a regra da aba Horas trabalhadas."
    )
    st.markdown(
        """
        <style>
        .tempo-barra {height:18px;background:#e2e8f0;border-radius:6px;overflow:hidden;
            display:flex;width:100%;box-shadow:inset 0 0 0 1px #cbd5e1;margin-top:4px}
        .tempo-segmento {height:18px;min-width:0}
        .tempo-vazia {background:#f8fafc}
        .tempo-valores {display:flex;gap:12px;flex-wrap:wrap;font-size:.72rem;
            color:#475569;margin-top:5px;margin-bottom:3px}
        .tempo-valores span {white-space:nowrap}
        .tempo-valores i {display:inline-block;width:9px;height:9px;border-radius:2px;
            margin-right:4px}
        .tempo-legenda {display:flex;gap:16px;flex-wrap:wrap;margin:8px 0 14px 0;
            color:#334155;font-size:.86rem}
        .tempo-legenda span {display:inline-flex;align-items:center;gap:5px}
        .tempo-legenda i {width:12px;height:12px;border-radius:3px;display:inline-block}
        .prioridade {display:inline-block;padding:2px 7px;border-radius:12px;color:white;
            font-size:.72rem;font-weight:700;margin-top:4px}
        .prioridade-1 {background:#991b1b}.prioridade-2 {background:#c2410c}
        .prioridade-3 {background:#a16207}.sem-prioridade {background:#64748b}
        .dia-linha {display:grid;grid-template-columns:120px minmax(420px,1fr) 120px;
            gap:14px;align-items:center;padding:8px 0;border-bottom:1px solid #e2e8f0}
        .dia-data {font-weight:700;color:#334155}.dia-total {font-size:.82rem;color:#334155}
        @media (max-width:900px) {.dia-linha {grid-template-columns:100px 1fr}
            .dia-total {grid-column:2}}
        </style>
        <div class="tempo-legenda">
          <span><i style="background:#16a34a"></i>Trabalhando</span>
          <span><i style="background:#2563eb"></i>Refeição</span>
          <span><i style="background:#dc2626"></i>Manutenção</span>
          <span><i style="background:#eab308"></i>Retorno à base</span>
          <span><i style="background:#94a3b8"></i>Outros intervalos</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    with st.expander("Como funcionam os alertas de prioridade?"):
        st.markdown(
            """
            - **Prioridade 1:** ao menos uma manutenção acima de 01:30, ou mais de
              uma manutenção de 01:00 ou mais no mês.
            - **Prioridade 2:** refeição acima de 01:15 em mais de dois dias no mês.
            - **Prioridade 3:** ao menos um retorno à base acima de 01:30.

            A ordem das equipes continua sendo definida pelo **total de intervalos**;
            a prioridade serve para destacar o tipo de ocorrência.
            """
        )

    painel_resumo, painel_dias, _ = montar_painel_tempos(
        dados, intervalos_individuais, ano, mes, grupos, equipes
    )
    if painel_resumo.empty:
        st.info("Nenhuma equipe disponível para os filtros selecionados.")
    else:
        quantidade_dias = calendar.monthrange(ano, mes)[1]
        hoje = datetime.now(FUSO_GOIAS).date()
        feriados = feriados_brasil(ano)
        for grupo in grupos:
            equipes_grupo = painel_resumo[painel_resumo["GRUPO"].eq(grupo)]
            if equipes_grupo.empty:
                continue
            with st.container(border=True):
                st.markdown(f"### {grupo}")
                for posicao, (_, item) in enumerate(equipes_grupo.iterrows(), 1):
                    coluna_equipe, coluna_barra, coluna_total = st.columns(
                        [2.0, 7.2, 1.8], vertical_alignment="center"
                    )
                    prioridade = int(item["PRIORIDADE"])
                    if prioridade:
                        selo = (
                            f'<span class="prioridade prioridade-{prioridade}">'
                            f'Prioridade {prioridade}</span>'
                        )
                    else:
                        selo = '<span class="prioridade sem-prioridade">Sem alerta</span>'
                    coluna_equipe.markdown(
                        f"**{posicao}. {item['PREFIXO']}**<br>{selo}",
                        unsafe_allow_html=True,
                    )
                    coluna_barra.markdown(
                        html_barra_tempos(
                            item["TRABALHO_HORAS"], item["REFEICAO_HORAS"],
                            item["MANUTENCAO_HORAS"], item["RETORNO_BASE_HORAS"],
                            item["OUTROS_INTERVALOS_HORAS"],
                        ),
                        unsafe_allow_html=True,
                    )
                    coluna_total.markdown(
                        "**Total de intervalos**<br>"
                        f"{formatar_duracao(item['TOTAL_INTERVALOS_HORAS'])}",
                        unsafe_allow_html=True,
                    )
                    if item["ALERTAS"]:
                        st.caption(f"Critério acionado — {item['PREFIXO']}: {item['ALERTAS']}")

                    with st.expander(f"Ver todos os dias de {item['PREFIXO']}"):
                        dias_equipe = painel_dias[
                            painel_dias["PREFIXO"].eq(item["PREFIXO"])
                        ]
                        linhas_dias_html = []
                        for dia in range(1, quantidade_dias + 1):
                            data_dia = date(ano, mes, dia)
                            registro = dias_equipe[dias_equipe["DATA"].eq(data_dia)]
                            if registro.empty:
                                if data_dia > hoje:
                                    situacao_dia = "Dia futuro"
                                elif data_dia in feriados:
                                    situacao_dia = "Feriado"
                                elif data_dia.weekday() == 5:
                                    situacao_dia = "Sábado"
                                elif data_dia.weekday() == 6:
                                    situacao_dia = "Domingo"
                                else:
                                    situacao_dia = "Sem turno"
                                barra_dia = (
                                    '<div class="tempo-barra tempo-vazia"></div>'
                                    f'<div class="tempo-valores">{situacao_dia}</div>'
                                )
                                total_dia = "—"
                            else:
                                dia_item = registro.iloc[0]
                                barra_dia = html_barra_tempos(
                                    dia_item["TRABALHO_HORAS"],
                                    dia_item["REFEICAO_HORAS"],
                                    dia_item["MANUTENCAO_HORAS"],
                                    dia_item["RETORNO_BASE_HORAS"],
                                    dia_item["OUTROS_INTERVALOS_HORAS"],
                                )
                                total_dia = formatar_duracao(
                                    dia_item["TOTAL_INTERVALOS_HORAS"]
                                )
                            linhas_dias_html.append(
                                '<div class="dia-linha">'
                                f'<div class="dia-data">{data_dia:%d/%m/%Y}</div>'
                                f'<div>{barra_dia}</div>'
                                f'<div class="dia-total"><strong>Intervalos</strong><br>{total_dia}</div>'
                                '</div>'
                            )
                        st.markdown(
                            '<div class="dia-lista">' + "".join(linhas_dias_html) + '</div>',
                            unsafe_allow_html=True,
                        )
                    st.divider()
with aba_mapa:
    st.subheader(f"Presença das equipes — {MESES[mes - 1]} de {ano}")
    st.caption(
        "🟩 1 = abriu turno  •  🟥 0 = não abriu em dia útil já transcorrido  •  "
        "S = sábado  •  D = domingo  •  F = feriado  •  vazio = dia futuro"
    )
    mapa = montar_mapa_mensal(dados, ano, mes, grupos, equipes)
    if mapa.empty:
        st.info("Nenhuma equipe disponível para os filtros selecionados.")
    else:
        colunas_dias = [coluna for coluna in mapa.columns if coluna != "TOTAL"]
        tabela_estilizada = (
            mapa.style
            .map(estilo_mapa, subset=colunas_dias)
            .set_properties(
                subset=["TOTAL"],
                **{"font-weight": "700", "background-color": "#f1f5f9"},
            )
        )
        st.dataframe(
            tabela_estilizada,
            use_container_width=True,
            height=min(800, 70 + len(mapa) * 35),
        )
with aba_horas:
    st.subheader(f"Horas trabalhadas — {MESES[mes - 1]} de {ano}")
    st.caption(
        "🟪 menos de 6 horas  •  🟥 de 6 horas até 07:59  •  "
        "🟩 8 horas ou mais  •  🟨 turno ainda em curso  •  "
        "- = nenhuma hora  •  S = sábado  •  D = domingo  •  F = feriado  •  "
        "vazio = dia futuro"
    )
    tabela_horas = montar_tabela_horas(dados, ano, mes, grupos, equipes)
    if tabela_horas.empty:
        st.info("Nenhuma equipe disponível para os filtros selecionados.")
    else:
        colunas_dias = [coluna for coluna in tabela_horas.columns if coluna != "TOTAL (h)"]
        horas_estilizadas = (
            tabela_horas.style
            .map(estilo_horas, subset=colunas_dias)
            .format(formatar_duracao, subset=colunas_dias + ["TOTAL (h)"])
            .map(estilo_total_horas, subset=["TOTAL (h)"])
        )
        st.dataframe(
            horas_estilizadas,
            use_container_width=True,
            height=min(800, 70 + len(tabela_horas) * 35),
        )
with aba_intervalos:
    st.subheader(f"Tempo de intervalo — {MESES[mes - 1]} de {ano}")
    st.caption(
        "Soma dos intervalos oficiais registrados para a equipe no dia.  "
        "🟧 menos de 01:00  •  🟩 de 01:00 até 01:15  •  "
        "🟥 de 01:16 até 02:30  •  "
        "🟪 acima de 02:30  •  "
        "branco = sem turno  •  "
        "S = sábado  •  D = domingo  •  F = feriado  •  vazio = dia futuro"
    )
    tabela_intervalos = montar_tabela_intervalos(dados, ano, mes, grupos, equipes)
    if tabela_intervalos.empty:
        st.info("Nenhuma equipe disponível para os filtros selecionados.")
    else:
        colunas_dias = [coluna for coluna in tabela_intervalos.columns if coluna != "TOTAL"]
        intervalos_estilizados = (
            tabela_intervalos.style
            .map(estilo_intervalos, subset=colunas_dias)
            .format(formatar_duracao, subset=colunas_dias + ["TOTAL"])
            .set_properties(
                subset=["TOTAL"],
                **{"font-weight": "700", "background-color": "#f1f5f9"},
            )
        )
        st.dataframe(
            intervalos_estilizados,
            use_container_width=True,
            height=min(800, 70 + len(tabela_intervalos) * 35),
        )
with aba_refeicao:
    st.subheader(f"Intervalos de refeição — {MESES[mes - 1]} de {ano}")
    st.caption(
        "Soma somente os intervalos oficiais classificados como refeição para a "
        "equipe no dia.  "
        "🟧 abaixo de 00:59  •  🟩 de 00:59 até 01:15  •  "
        "🟥 acima de 01:15 até 01:39  •  🟪 acima de 01:39  •  "
        "branco = sem turno  •  "
        "S = sábado  •  D = domingo  •  F = feriado  •  vazio = dia futuro"
    )
    tabela_refeicao = montar_tabela_refeicao(
        dados, intervalos_individuais, ano, mes, grupos, equipes
    )
    if tabela_refeicao.empty:
        st.info("Nenhuma equipe disponível para os filtros selecionados.")
    else:
        colunas_dias = [coluna for coluna in tabela_refeicao.columns if coluna != "TOTAL"]
        refeicao_estilizada = (
            tabela_refeicao.style
            .map(estilo_refeicao, subset=colunas_dias)
            .format(formatar_duracao, subset=colunas_dias + ["TOTAL"])
            .set_properties(
                subset=["TOTAL"],
                **{"font-weight": "700", "background-color": "#f1f5f9"},
            )
        )
        st.dataframe(
            refeicao_estilizada,
            use_container_width=True,
            height=min(800, 70 + len(tabela_refeicao) * 35),
        )
with aba_ranking:
    st.subheader(f"Maiores intervalos por motivo — {MESES[mes - 1]} de {ano}")
    st.caption(
        "Cada linha representa um intervalo oficial individual. Os valores não são "
        "somados por equipe nem por dia. Cada tabela possui sua própria classificação."
    )
    if ranking_intervalos.empty:
        st.info(
            "Nenhum intervalo individual encerrado foi encontrado. Execute novamente "
            "o bot atualizado para enviar os detalhes dos intervalos."
        )
    else:
        categorias_ranking = [
            ("Refeição", "REFEICAO", "refeicao"),
            ("Manutenção no veículo", "MANUTENCAO_NO_VEICULO", "manutencao"),
            ("Retorno para a base", "RETORNO_PARA_BASE", "retorno_base"),
        ]
        for titulo_categoria, codigo_categoria, chave_categoria in categorias_ranking:
            st.markdown(f"#### {titulo_categoria}")
            tabela_categoria = ranking_intervalos[
                ranking_intervalos["CATEGORIA_MOTIVO"].eq(codigo_categoria)
            ].copy().reset_index(drop=True)
            tabela_categoria.insert(
                0, "POSICAO", range(1, len(tabela_categoria) + 1)
            )
            if tabela_categoria.empty:
                st.info(f"Nenhum intervalo de {titulo_categoria.lower()} no período.")
                continue

            colunas_ranking = [
                "POSICAO", "PREFIXO", "INICIO_INTERVALO", "FIM_INTERVALO",
                "DURACAO_INTERVALO", "MOTIVO_INTERVALO",
            ]
            st.dataframe(
                tabela_categoria[colunas_ranking],
                hide_index=True,
                use_container_width=True,
                height=min(600, 70 + len(tabela_categoria) * 35),
                column_config={
                    "POSICAO": st.column_config.NumberColumn("Posição", format="%d"),
                    "PREFIXO": st.column_config.TextColumn("Equipe"),
                    "INICIO_INTERVALO": st.column_config.DatetimeColumn(
                        "Início do intervalo", format="DD/MM/YYYY HH:mm"
                    ),
                    "FIM_INTERVALO": st.column_config.DatetimeColumn(
                        "Fim do intervalo", format="DD/MM/YYYY HH:mm"
                    ),
                    "DURACAO_INTERVALO": st.column_config.TextColumn("Duração"),
                    "MOTIVO_INTERVALO": st.column_config.TextColumn("Motivo"),
                },
            )
            st.download_button(
                f"Baixar ranking de {titulo_categoria.lower()} em CSV",
                tabela_categoria[colunas_ranking].to_csv(
                    index=False, sep=";", decimal=","
                ).encode("utf-8-sig"),
                file_name=(
                    f"ranking_{chave_categoria}_{ano}_{mes:02}.csv"
                ),
                mime="text/csv",
                key=f"baixar_ranking_{chave_categoria}",
            )
with aba_resumo:
    if filtrado.empty:
        st.info("Nenhum turno encontrado para os filtros selecionados.")
    else:
        resumo = filtrado.pivot_table(
            index="DATA", columns="GRUPO", values="HIST_TURMA_PLANTAO_ID",
            aggfunc="nunique", fill_value=0,
        ).reindex(columns=GRUPOS, fill_value=0)
        resumo["TOTAL"] = resumo.sum(axis=1)
        st.dataframe(resumo, use_container_width=True)
        if grupos:
            st.line_chart(resumo[grupos])

st.download_button(
    "Baixar tabela filtrada em CSV",
    filtrado[colunas].to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
    file_name=f"acompanhamento_turnos_{ano}_{mes:02}.csv",
    mime="text/csv",
)
